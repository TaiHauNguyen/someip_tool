#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dbc_bitfield_excel.py  (ban va loi 2026-09)
==========================================
Sinh BITFIELD STRUCT tu file CAN .dbc va ghi vao sheet ``DataStructures``.

Thay doi so voi ban cu
----------------------
1. DO COT LINH HOAT (sua loi "Thieu cot 'enumerate' trong dong tieu de 2")
   * so khop ten cot sau khi chuan hoa (bo dau cach, '_', '-', '.', chu hoa)
   * chap nhan bi danh: enumerate / enum / enumeration / enumerator /
     literal / label / element / member ...
   * cell gop (merged) lay gia tri o o goc trai
   * dong tieu de duoc do trong 3 dong ke duoi dong title, chon dong khop nhieu nhat
   * neu van thieu -> in ra TOAN BO dong tieu de doc duoc de biet cot that ten gi,
     va co the chi dinh tay bang --enum-cols / --struct-cols
2. DATATYPE 1:1 VOI DBC
   * co dau: signal '-' trong DBC -> intN_t (truoc day luon uintN_t)
   * kieu nen = max(kieu theo do dai bit, kieu theo so byte ma field trai qua)
   * enum: kieu nen lay theo do dai + dau cua signal (truoc day hard-code uint8_t)
   * gia tri enum am ghi dang thap phan (truoc day '0x-1')
   * signal float/double (SIG_VALTYPE_) khong ghi ':' bit vi C khong cho
     bitfield kieu float -> canh bao ro rang
   * kiem tra gia tri VAL_ co vuot ngoai do rong bit hay khong

Usage
-----
  python dbc_bitfield_excel.py --excel ACM_1.2_split.xlsx --dbc X.dbc --dry-run
  python dbc_bitfield_excel.py --excel ACM_1.2_split.xlsx --dbc X.dbc --dump-header
  python dbc_bitfield_excel.py --excel ACM_1.2_split.xlsx --dbc X.dbc \
      --enum-cols A,B,C,D,E,F --struct-cols H,I,J,K,L

Requires: openpyxl
"""

from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
from copy import copy

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Can openpyxl: python -m pip install openpyxl")

from dbc_struct import parse_dbc, build_layout

SHEET_NAME = "DataStructures"


# --------------------------------------------------------------------------- #
# Naming
# --------------------------------------------------------------------------- #
def camel(text, upper_first=True):
    """``No_Airbag_Deployment`` -> ``NoAirbagDeployment``."""
    parts = [p for p in re.split(r"[^0-9A-Za-z]+", str(text)) if p]
    if not parts:
        return ""
    out = []
    for i, part in enumerate(parts):
        if i == 0 and not upper_first:
            out.append(part[0].lower() + part[1:])
        else:
            out.append(part[0].upper() + part[1:])
    return "".join(out)


def enum_type_name(signal_name):
    """``ACU_Airbag_Deployment_Status`` -> ``ACUAirbagDeploymentStatusType``."""
    return camel(signal_name) + "Type"


def struct_type_name(message_name):
    """``ACU_CRASH_INFO`` -> ``ACUCrashInfoStruct``."""
    parts = [p for p in re.split(r"[^0-9A-Za-z]+", message_name) if p]
    out = []
    for part in parts:
        if part.isupper() and len(part) > 3:
            out.append(part[0] + part[1:].lower())
        elif part.isupper():
            out.append(part)
        else:
            out.append(part[0].upper() + part[1:])
    return "".join(out) + "Struct"


# --------------------------------------------------------------------------- #
# Datatype: map 1:1 voi DBC
# --------------------------------------------------------------------------- #
def _int_ctype(nbits, signed):
    return ("int%d_t" % nbits) if signed else ("uint%d_t" % nbits)


def _storage_bits(length, byte_span):
    """So bit cua kieu nen: du chua do dai bit VA du chua so byte field trai qua."""
    need = max(length, byte_span * 8)
    for n in (8, 16, 32, 64):
        if need <= n:
            return n
    return 64


def field_byte_span(field):
    first = field.wire_start // 8
    last = (field.wire_start + field.width - 1) // 8
    return last - first + 1


def elem_ctype(field):
    """Kieu nen cua mot element trong bitfield struct, giu dung dau cua DBC."""
    sig = field.signal
    if sig is not None and getattr(sig, "is_float", False):
        return "float" if sig.length == 32 else "double"
    signed = bool(sig is not None and sig.is_signed)
    return _int_ctype(_storage_bits(field.width, field_byte_span(field)), signed)


def enum_base_ctype(sig):
    """Kieu nen cua enum = theo do dai + dau cua chinh signal trong DBC."""
    return _int_ctype(_storage_bits(sig.length, 0), bool(sig.is_signed))


def fmt_enum_value(v, sig):
    """Gia tri enum: hex cho so khong am, thap phan cho so am."""
    if v < 0:
        return "%d" % v
    return "0x%X" % v


def check_enum_range(sig, warn):
    """Canh bao neu gia tri trong VAL_ khong lot vao do rong bit cua signal."""
    if not sig.values:
        return
    if sig.is_signed:
        lo, hi = -(1 << (sig.length - 1)), (1 << (sig.length - 1)) - 1
    else:
        lo, hi = 0, (1 << sig.length) - 1
    bad = [v for v in sig.values if v < lo or v > hi]
    if bad:
        warn("enum '%s': gia tri %s nam ngoai %d bit (%d..%d)"
             % (sig.name, ", ".join(str(b) for b in sorted(bad)), sig.length, lo, hi))


# --------------------------------------------------------------------------- #
# Do vi tri cot cua sheet
# --------------------------------------------------------------------------- #
def _norm(v):
    """'Enum. Name ' -> 'enumname'  (bo moi ky tu khong phai chu/so)."""
    if v is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", str(v).strip().lower())


# Bi danh cho tung cot. Xep tu cu the -> chung chung.
ENUM_ALIASES = {
    "index":       ["index", "idx", "no", "stt", "num", "id"],
    "name":        ["name", "typename", "enumtypename", "datatypename", "enumname"],
    "type":        ["type", "basetype", "datatype", "basedatatype", "underlyingtype"],
    "enumerate":   ["enumerate", "enum", "enumeration", "enumerator", "enumelement",
                    "enumliteral", "enumitem", "literal", "label", "member",
                    "element", "enumvaluename", "constant"],
    "value":       ["value", "val", "enumvalue", "code", "rawvalue"],
    "description": ["description", "desc", "comment", "note", "remark", "meaning"],
}
STRUCT_ALIASES = {
    "index":       ["index", "idx", "no", "stt", "num", "id"],
    "name":        ["name", "structname", "typename", "datatypename"],
    "element":     ["element", "member", "field", "elementname", "membername",
                    "fieldname", "signal", "signalname"],
    "type":        ["type", "datatype", "elementtype", "membertype", "basetype"],
    "description": ["description", "desc", "comment", "note", "remark", "meaning"],
}
# Thu tu cot mac dinh khi phai doan theo vi tri.
ENUM_ORDER = ["index", "name", "type", "enumerate", "value", "description"]
STRUCT_ORDER = ["index", "name", "element", "type", "description"]


class SheetError(Exception):
    pass


class SheetMap(object):
    """Vi tri cac cot cua khoi Enumerate va khoi Struct."""

    def __init__(self, ws, enum_cols=None, struct_cols=None, verbose=True):
        self.ws = ws
        self.enum = {}
        self.struct = {}
        self.title_row = None
        self.header_row = None
        self.notes = []
        self._merged_lookup = self._build_merged_lookup(ws)
        if enum_cols or struct_cols:
            self._manual(enum_cols, struct_cols)
        else:
            self._detect()
        self.first_data_row = self.header_row + 1
        if verbose:
            for n in self.notes:
                sys.stderr.write("NOTE: %s\n" % n)

    # ---- doc cell, ho tro merged cell ------------------------------------ #
    @staticmethod
    def _build_merged_lookup(ws):
        look = {}
        merged = getattr(ws, "merged_cells", None)
        ranges = list(merged.ranges) if merged is not None else []
        for rng in ranges:
            top = (rng.min_row, rng.min_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    look[(r, c)] = top
        return look

    def cell_value(self, row, col):
        r, c = self._merged_lookup.get((row, col), (row, col))
        return self.ws.cell(r, c).value

    # ---- chi dinh tay ---------------------------------------------------- #
    def _manual(self, enum_cols, struct_cols):
        def to_idx(spec, order, what):
            cols = [c.strip() for c in spec.split(",") if c.strip()]
            if len(cols) != len(order):
                raise SheetError("--%s can dung %d cot (%s), nhan duoc %d"
                                 % (what, len(order), ", ".join(order), len(cols)))
            out = {}
            for key, c in zip(order, cols):
                out[key] = int(c) if c.isdigit() else column_index_from_string(c.upper())
            return out

        if not (enum_cols and struct_cols):
            raise SheetError("Phai chi dinh CA --enum-cols va --struct-cols")
        self.enum = to_idx(enum_cols, ENUM_ORDER, "enum-cols")
        self.struct = to_idx(struct_cols, STRUCT_ORDER, "struct-cols")
        self.title_row = 1
        self.header_row = self._guess_header_row_manual()
        self.notes.append("Dung cot do nguoi dung chi dinh, dong tieu de = %d"
                          % self.header_row)

    def _guess_header_row_manual(self):
        want = _norm("index")
        for row in range(1, min(self.ws.max_row, 10) + 1):
            if _norm(self.cell_value(row, self.enum["index"])) == want:
                return row
        return 2

    # ---- do tu dong ------------------------------------------------------ #
    def _detect(self):
        ws = self.ws
        max_scan_row = min(ws.max_row, 12)
        max_col = ws.max_column

        enum_col = struct_col = None
        for row in range(1, max_scan_row + 1):
            for col in range(1, max_col + 1):
                key = _norm(self.cell_value(row, col))
                if key in ("enumerate", "enum", "enumeration") and enum_col is None:
                    enum_col, self.title_row = col, row
                elif key in ("struct", "structs", "structure", "structures") and struct_col is None:
                    struct_col = col
            if enum_col is not None and struct_col is not None:
                break
        if enum_col is None or struct_col is None:
            raise SheetError(
                "Khong tim thay title 'Enumerate' / 'Struct' o dau sheet.\n"
                + self.dump_rows(1, min(6, ws.max_row)))
        if struct_col <= enum_col:
            raise SheetError("Cot 'Struct' (%s) nam truoc cot 'Enumerate' (%s) - khong ho tro"
                             % (get_column_letter(struct_col), get_column_letter(enum_col)))

        # dong tieu de: thu title_row+1 .. title_row+3, chon dong khop nhieu nhat
        best = None
        for hrow in range(self.title_row + 1, min(self.title_row + 4, ws.max_row) + 1):
            hdr = self._header_of(hrow, max_col)
            e_map, e_miss = self._match(hdr, enum_col, struct_col, ENUM_ALIASES, ENUM_ORDER)
            s_map, s_miss = self._match(hdr, struct_col, max_col + 1, STRUCT_ALIASES, STRUCT_ORDER)
            score = (len(ENUM_ORDER) - len(e_miss)) + (len(STRUCT_ORDER) - len(s_miss))
            if best is None or score > best[0]:
                best = (score, hrow, hdr, e_map, e_miss, s_map, s_miss)
            if not e_miss and not s_miss:
                break

        score, hrow, hdr, e_map, e_miss, s_map, s_miss = best
        self.header_row = hrow

        # con thieu -> doan theo vi tri trong khoi (co canh bao)
        if e_miss:
            e_map = self._fill_positional(e_map, e_miss, enum_col, struct_col,
                                          ENUM_ORDER, "Enumerate", hdr)
        if s_miss:
            s_map = self._fill_positional(s_map, s_miss, struct_col, max_col + 1,
                                          STRUCT_ORDER, "Struct", hdr)
        self.enum, self.struct = e_map, s_map

    def _header_of(self, row, max_col):
        hdr = {}
        for col in range(1, max_col + 1):
            n = _norm(self.cell_value(row, col))
            if n:
                hdr[col] = n
        return hdr

    @staticmethod
    def _match(hdr, start, end, aliases, order):
        """Gan cot cho tung key: khop chinh xac truoc, roi den khop chua/bat dau."""
        cols = [c for c in range(start, end) if c in hdr]
        used = set()
        got = {}
        # vong 1: khop chinh xac voi bi danh
        for key in order:
            for c in cols:
                if c in used:
                    continue
                if hdr[c] in aliases[key]:
                    got[key] = c
                    used.add(c)
                    break
        # vong 2: khop "bat dau bang" / "co chua"
        for key in order:
            if key in got:
                continue
            for c in cols:
                if c in used:
                    continue
                h = hdr[c]
                if any(h.startswith(a) or a.startswith(h) or a in h for a in aliases[key]):
                    got[key] = c
                    used.add(c)
                    break
        missing = [k for k in order if k not in got]
        return got, missing

    def _fill_positional(self, got, missing, start, end, order, block, hdr):
        """Doan cot con thieu theo vi tri chuan cua khoi."""
        span = [c for c in range(start, min(end, start + len(order)))]
        for key in missing:
            guess = span[order.index(key)] if order.index(key) < len(span) else None
            if guess is None or guess in got.values():
                raise SheetError(
                    "Khoi %s: khong xac dinh duoc cot '%s' o dong tieu de %d.\n"
                    "Doc duoc: %s\n"
                    "Cach xu ly: sua ten cot trong Excel, hoac chay lai voi\n"
                    "  --enum-cols A,B,C,D,E,F --struct-cols H,I,J,K,L\n"
                    "(dung thu tu: enum = %s ; struct = %s)"
                    % (block, key, self.header_row,
                       ", ".join("%s='%s'" % (get_column_letter(c), hdr[c])
                                 for c in sorted(hdr) if start <= c < end) or "(trong)",
                       "|".join(ENUM_ORDER), "|".join(STRUCT_ORDER)))
            got[key] = guess
            self.notes.append("Khoi %s: khong thay tieu de '%s', tam dung cot %s theo vi tri"
                              % (block, key, get_column_letter(guess)))
        return got

    # ---- chan doan ------------------------------------------------------- #
    def dump_rows(self, r0, r1):
        out = ["--- Noi dung doc duoc tu sheet ---"]
        for row in range(r0, r1 + 1):
            cells = []
            for col in range(1, min(self.ws.max_column, 30) + 1):
                v = self.cell_value(row, col)
                if v is not None and str(v).strip():
                    cells.append("%s%d='%s'" % (get_column_letter(col), row, str(v).strip()))
            out.append("  row %d: %s" % (row, "  ".join(cells) if cells else "(trong)"))
        return "\n".join(out)

    def describe(self):
        e = ", ".join("%s=%s" % (k, get_column_letter(self.enum[k])) for k in ENUM_ORDER)
        s = ", ".join("%s=%s" % (k, get_column_letter(self.struct[k])) for k in STRUCT_ORDER)
        return ("Title row %d, header row %d\n  Enumerate: %s\n  Struct   : %s"
                % (self.title_row, self.header_row, e, s))


# --------------------------------------------------------------------------- #
# Xay du lieu tu DBC
# --------------------------------------------------------------------------- #
def build_enums(messages, warn):
    """{enum_type: [(literal, value, description), ...]} + kieu nen theo DBC."""
    enums = {}
    base = {}
    order = []
    for msg in messages:
        for sig in msg.signals:
            if not sig.values:
                continue
            check_enum_range(sig, warn)
            name = enum_type_name(sig.name)
            items = [(camel(sig.values[v]), fmt_enum_value(v, sig), sig.values[v])
                     for v in sorted(sig.values)]
            if name in enums:
                if enums[name] != items:
                    warn("enum '%s' co 2 bang gia tri khac nhau, giu ban dau" % name)
                elif base[name] != enum_base_ctype(sig):
                    warn("enum '%s' co 2 kieu nen khac nhau (%s vs %s), giu ban dau"
                         % (name, base[name], enum_base_ctype(sig)))
                continue
            enums[name] = items
            base[name] = enum_base_ctype(sig)
            order.append(name)
    return [(n, base[n], enums[n]) for n in order]


def build_struct_rows(msg, include_reserved=True, warn=None):
    """(element, type_str, description, width_bits) cho bitfield struct cua 1 message."""
    warn = warn or (lambda m: None)
    rows = []
    for f in build_layout(msg):
        b0 = f.wire_start // 8
        p0 = 7 - (f.wire_start % 8)
        p1 = p0 - f.width + 1
        pos = ("byte%d [%d:%d]" % (b0, p0, p1) if p1 >= 0
               else "byte%d..byte%d" % (b0, (f.wire_start + f.width - 1) // 8))

        if f.is_reserved:
            if not include_reserved:
                continue
            rows.append((f.name,
                         "%s : %d" % (elem_ctype(f), f.width),
                         "padding; %s; wire bit %d..%d"
                         % (pos, f.wire_start, f.wire_start + f.width - 1),
                         f.width))
            continue

        sig = f.signal
        is_float = bool(getattr(sig, "is_float", False))
        if sig.values:
            base = enum_type_name(sig.name)
        else:
            base = elem_ctype(f)

        if is_float:
            # C khong cho bitfield kieu float/double -> ghi kieu tran, khong co ':'
            type_str = elem_ctype(f)
            warn("signal '%s' la float %d bit: C khong ho tro bitfield float, "
                 "element duoc ghi khong co ': %d'" % (sig.name, sig.length, sig.length))
        else:
            type_str = "%s : %d" % (base, f.width)

        desc = ["%s.%s" % (msg.name, sig.name)]
        if sig.comment:
            desc.append(sig.comment)
        desc.append(pos)
        desc.append("dbc start=%d, len=%d, %s%s"
                    % (sig.start_bit, sig.length,
                       "Motorola" if sig.is_motorola else "Intel",
                       ", signed" if sig.is_signed else ", unsigned"))
        if sig.values:
            desc.append("enum base=%s" % enum_base_ctype(sig))
        if sig.unit:
            desc.append("unit=%s" % sig.unit)
        if sig.factor != 1 or sig.offset != 0:
            desc.append("phys=raw*%g%+g (element luu RAW)" % (sig.factor, sig.offset))
        rows.append((f.name, type_str, "; ".join(desc), f.width))
    return rows


def structs_already_in_sheet(ws, smap):
    """Doc ten CAN message dang duoc tham chieu trong sheet, giu thu tu."""
    names = []
    col_desc = smap.struct["description"]
    for row in range(smap.first_data_row, ws.max_row + 1):
        v = ws.cell(row, col_desc).value
        if isinstance(v, str) and "." in v:
            m = re.match(r"\s*([A-Za-z_]\w*)\s*\.", v)
            if m and m.group(1) not in names:
                names.append(m.group(1))
    return names


# --------------------------------------------------------------------------- #
# Ghi sheet
# --------------------------------------------------------------------------- #
def _clear_block(ws, cols, first_row):
    for row in range(first_row, ws.max_row + 1):
        for col in cols:
            ws.cell(row, col).value = None


def _styles(ws, cols, row):
    return {c: copy(ws.cell(row, c)._style) for c in cols}


def write_sheet(ws, smap, enums, structs):
    enum_cols = [smap.enum[k] for k in ENUM_ORDER]
    struct_cols = [smap.struct[k] for k in STRUCT_ORDER]
    first = smap.first_data_row

    enum_style = _styles(ws, enum_cols, first)
    struct_style = _styles(ws, struct_cols, first)
    _clear_block(ws, enum_cols, first)
    _clear_block(ws, struct_cols, first)

    def put(row, col, value, styles):
        cell = ws.cell(row, col)
        cell._style = copy(styles[col])
        cell.value = value

    # ---- khoi Enumerate
    row = first
    for idx, (type_name, base_type, items) in enumerate(enums, start=1):
        for i, (literal, value, desc) in enumerate(items):
            if i == 0:
                put(row, smap.enum["index"], idx, enum_style)
                put(row, smap.enum["name"], type_name, enum_style)
                put(row, smap.enum["type"], base_type, enum_style)
            put(row, smap.enum["enumerate"], literal, enum_style)
            put(row, smap.enum["value"], value, enum_style)
            put(row, smap.enum["description"], desc, enum_style)
            row += 1
    enum_rows = row - first

    # ---- khoi Struct
    row = first
    for idx, (struct_name, rows) in enumerate(structs, start=1):
        for i, item in enumerate(rows):
            element, ctype, desc = item[0], item[1], item[2]
            if i == 0:
                put(row, smap.struct["index"], idx, struct_style)
                put(row, smap.struct["name"], struct_name, struct_style)
            put(row, smap.struct["element"], element, struct_style)
            put(row, smap.struct["type"], ctype, struct_style)
            put(row, smap.struct["description"], desc, struct_style)
            row += 1
    struct_rows = row - first
    return enum_rows, struct_rows


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Sinh bitfield struct tu DBC va ghi vao sheet DataStructures")
    ap.add_argument("--excel", required=True, help="workbook SOME/IP (.xlsx)")
    ap.add_argument("--dbc", required=True, nargs="+", help="mot hoac nhieu file .dbc")
    ap.add_argument("--msg", nargs="*", default=None,
                    help="ten CAN message; mac dinh lay dung bo dang co trong sheet")
    ap.add_argument("--sheet", default=SHEET_NAME)
    ap.add_argument("--out", help="file .xlsx dau ra (mac dinh: <ten>_bitfield.xlsx)")
    ap.add_argument("--inplace", action="store_true", help="ghi de, tao ban .bak")
    ap.add_argument("--dry-run", action="store_true", help="chi in, khong ghi file")
    ap.add_argument("--no-reserved", action="store_true",
                    help="bo cac element Reserved_n (struct se KHONG con 1:1)")
    ap.add_argument("--dump-header", action="store_true",
                    help="in 8 dong dau cua sheet roi thoat (chan doan cot)")
    ap.add_argument("--enum-cols",
                    help="chi dinh tay 6 cot khoi Enumerate, vd: A,B,C,D,E,F")
    ap.add_argument("--struct-cols",
                    help="chi dinh tay 5 cot khoi Struct, vd: H,I,J,K,L")
    a = ap.parse_args(argv)

    warnings = []

    def warn(m):
        warnings.append(m)
        sys.stderr.write("WARN: %s\n" % m)

    wb = openpyxl.load_workbook(a.excel)
    if a.sheet not in wb.sheetnames:
        sys.exit("Workbook khong co sheet '%s'. Sheet dang co: %s"
                 % (a.sheet, ", ".join(wb.sheetnames)))
    ws = wb[a.sheet]

    if a.dump_header:
        probe = SheetMap.__new__(SheetMap)
        probe.ws = ws
        probe._merged_lookup = SheetMap._build_merged_lookup(ws)
        print(probe.dump_rows(1, min(8, ws.max_row)))
        return 0

    try:
        smap = SheetMap(ws, a.enum_cols, a.struct_cols)
    except SheetError as e:
        sys.stderr.write("LOI DO COT:\n%s\n" % e)
        return 2

    by_name = {}
    for path in a.dbc:
        for m in parse_dbc(path):
            by_name.setdefault(m.name, m)

    wanted = a.msg if a.msg else structs_already_in_sheet(ws, smap)
    if not wanted:
        sys.exit("Khong xac dinh duoc message nao. Dung --msg de chi dinh.")

    missing = [n for n in wanted if n not in by_name]
    if missing:
        sys.exit("Khong co trong DBC: %s" % ", ".join(missing))

    messages = [by_name[n] for n in wanted]
    enums = build_enums(messages, warn)
    structs = []
    for m in messages:
        try:
            structs.append((struct_type_name(m.name),
                            build_struct_rows(m, include_reserved=not a.no_reserved,
                                              warn=warn)))
        except ValueError as e:
            sys.stderr.write(
                "LOI LAYOUT o message '%s': %s\n"
                "(thuong gap khi message co signal multiplex m0/m1 hoac signal chong bit; "
                "bitfield 1:1 khong bieu dien duoc truong hop nay)\n" % (m.name, e))
            return 4

    print("Sheet : %s" % a.sheet)
    print(smap.describe())
    print("Dong du lieu tu : %d" % smap.first_data_row)
    print("Message : %s" % ", ".join(wanted))
    print()

    bad_dlc = 0
    for m, (sname, rows) in zip(messages, structs):
        total = sum(r[3] for r in rows)
        ok = (total == m.dlc * 8)
        if not ok:
            bad_dlc += 1
        print("%-26s %-24s %2d element, %3d bit / %d byte %s"
              % (m.name, sname, len(rows), total, m.dlc,
                 "OK" if ok else "!! KHONG KHOP DLC"))
        for element, ctype, _desc, _w in rows:
            print("   %-36s %s" % (element, ctype))
        print()

    print("Enumerate: %d type" % len(enums))
    if warnings:
        print("Canh bao : %d (xem stderr)" % len(warnings))

    if a.dry_run:
        print("\n--dry-run: khong ghi file.")
        return 0
    if bad_dlc:
        sys.stderr.write("Co %d message khong khop DLC -> khong ghi file.\n" % bad_dlc)
        return 3

    enum_rows, struct_rows = write_sheet(ws, smap, enums, structs)

    if a.inplace:
        out = a.excel
        bak = a.excel + ".bak"
        shutil.copy2(a.excel, bak)
        print("\nBan sao luu -> %s" % bak)
    else:
        out = a.out or (os.path.splitext(a.excel)[0] + "_bitfield.xlsx")

    wb.save(out)
    print("Da ghi %d dong enum, %d dong struct -> %s" % (enum_rows, struct_rows, out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
